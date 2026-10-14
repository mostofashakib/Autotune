"""Document ingestion keeps provenance and stops on incomplete generation."""
import json
import zipfile
from unittest.mock import Mock

import pytest

from autotune.core.domain.models import ModelEndpoint, TaskType


def ingestor(replies, **kwargs):
    from autotune.infrastructure.sanitizers.document_ingestor import DocumentIngestor, DocumentOptions
    client = Mock()
    client.complete.side_effect = replies
    return DocumentIngestor(
        ModelEndpoint(base_url="http://localhost:11434/v1", model="local"),
        client=client, options=DocumentOptions(**kwargs),
    )


def pair(topic):
    return json.dumps([{"question": f"Describe the {topic} policy", "answer": f"The {topic} policy requires daily review."}])


def test_documents_are_split_by_source_and_report_progress(tmp_path):
    archive = tmp_path / "docs.zip"
    with zipfile.ZipFile(archive, "w") as z:
        for name in ("safety", "billing", "access"):
            z.writestr(f"{name}.md", f"The {name} policy requires daily review.")
    progress = []
    worker = ingestor([item for name in ("access", "billing", "safety") for item in (pair(name), '{"supported": true}')])
    worker.on_progress = progress.append
    result = worker.sanitize(archive)
    assert result.task_type == TaskType.CHAT
    assert result.total_records == 3
    train = {r.metadata["source_file"] for r in result.train_records}
    val = {r.metadata["source_file"] for r in result.val_records}
    assert train and val and train.isdisjoint(val)
    assert all(r.metadata["chunk_id"] for r in result.train_records + result.val_records)
    assert progress[-1] == {"documents_read": 3, "chunks_made": 3, "pairs_generated": 3, "pairs_kept": 3}


def test_failed_generator_never_returns_partial_data(tmp_path):
    from autotune.infrastructure.evaluators.openai_client import ModelCallError
    path = tmp_path / "data.md"
    path.write_text("The safety policy requires daily review.")
    with pytest.raises(ModelCallError, match="offline"):
        ingestor([ModelCallError("offline")]).sanitize(path)


@pytest.mark.parametrize("reply", [
    '[{"question":"Describe the safety policy", "answer":"Martian dragons fly above purple oceans."}]',
    '[{"question":"Why", "answer":"No"}]',
])
def test_unsupported_or_short_pairs_are_not_training_data(tmp_path, reply):
    path = tmp_path / "data.txt"
    path.write_text("The safety policy requires daily review.")
    with pytest.raises(ValueError, match="No grounded"):
        ingestor([reply]).sanitize(path)


def test_grounding_check_rejects_plausible_but_false_answer(tmp_path):
    path = tmp_path / "data.md"
    path.write_text("The safety policy requires daily review.")
    with pytest.raises(ValueError, match="No grounded"):
        ingestor([pair("safety"), '{"supported": false}']).sanitize(path)


def test_duplicates_and_near_duplicates_are_dropped(tmp_path):
    path = tmp_path / "data.md"
    path.write_text("The safety policy requires daily review.")
    row = json.loads(pair("safety"))[0]
    reply = json.dumps([row, row, {**row, "question": row["question"] + " please"}])
    result = ingestor([reply, '{"supported": true}']).sanitize(path)
    assert result.total_records == 1
    assert result.dropped_count == 2
    assert "one document" in result.warnings[0]


def test_chunk_cap_fails_instead_of_truncating(tmp_path):
    path = tmp_path / "data.md"
    path.write_text("word " * 100)
    with pytest.raises(ValueError, match="chunk cap"):
        ingestor([], chunk_tokens=10, overlap_tokens=2, max_chunks=2).sanitize(path)


def test_zip_paths_are_not_extracted(tmp_path):
    path = tmp_path / "data.zip"
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("../escape.md", "safety policy")
    with pytest.raises(ValueError, match="Unsafe"):
        ingestor([]).sanitize(path)
    assert not (tmp_path.parent / "escape.md").exists()


def test_missing_model_names_configuration(monkeypatch):
    from autotune.settings import document_endpoint
    monkeypatch.setenv("AUTOTUNE_DOCUMENT_MODEL", "")
    with pytest.raises(ValueError, match="AUTOTUNE_DOCUMENT_MODEL"):
        document_endpoint()


def test_token_chunks_have_overlap():
    from autotune.infrastructure.sanitizers.document_ingestor import chunk_text
    # Injected tokenizer avoids network access and makes token boundaries explicit.
    class Tokens:
        def encode(self, text): return text.split()
        def decode(self, tokens): return " ".join(tokens)
    assert chunk_text("a b c d e f", 4, 1, Tokens()) == ["a b c d", "d e f"]


def test_html_ignores_scripts_and_extracts_content():
    from autotune.infrastructure.sanitizers.document_ingestor import extract_text
    assert extract_text('page.html', b'<style>hidden</style><p>Public policy</p><script>secret</script>') == 'Public policy'


def test_word_extracts_paragraphs_and_tables():
    import io
    from docx import Document
    from autotune.infrastructure.sanitizers.document_ingestor import extract_text
    document = Document()
    document.add_paragraph('Public policy')
    document.add_table(rows=1, cols=1).cell(0, 0).text = 'Daily review'
    buffer = io.BytesIO()
    document.save(buffer)
    assert extract_text('policy.docx', buffer.getvalue()) == 'Public policy\nDaily review'


def test_blank_pdf_needs_ocr():
    import io
    from pypdf import PdfWriter
    from autotune.infrastructure.sanitizers.document_ingestor import extract_text
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ValueError, match='OCR'):
        extract_text('scan.pdf', buffer.getvalue())


def test_missing_extras_names_install_command(monkeypatch):
    import sys
    from autotune.infrastructure.sanitizers.document_ingestor import extract_text
    monkeypatch.setitem(sys.modules, 'pypdf', None)
    with pytest.raises(RuntimeError, match='uv sync --extra documents'):
        extract_text('policy.txt', b'policy')


@pytest.mark.parametrize('verdict', ['{"supported": "true"}', 'not json'])
def test_invalid_grounding_reply_fails_whole_upload(tmp_path, verdict):
    from autotune.infrastructure.evaluators.openai_client import ModelCallError
    path = tmp_path / 'policy.md'
    path.write_text('The safety policy requires daily review.')
    with pytest.raises(ModelCallError):
        ingestor([pair('safety'), verdict]).sanitize(path)


def test_document_key_placeholder_is_rejected(monkeypatch):
    from autotune.settings import document_endpoint, CredentialNotConfigured
    monkeypatch.setenv('AUTOTUNE_DOCUMENT_MODEL', 'local')
    monkeypatch.setenv('AUTOTUNE_DOCUMENT_API_KEY_ENV', 'DOC_TEST_KEY')
    monkeypatch.setenv('DOC_TEST_KEY', 'mock_key')
    with pytest.raises(CredentialNotConfigured, match='DOC_TEST_KEY'):
        document_endpoint()
